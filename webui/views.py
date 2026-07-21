import io
import json

from datetime import date, datetime

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.views.decorators.http import require_POST, require_http_methods
from django.contrib import messages
from django.utils import timezone
from django.urls import reverse
from django.core.paginator import Paginator, EmptyPage
from django.db.models import Q, Count, Case, When, IntegerField
from django.http import JsonResponse
from openpyxl.utils import get_column_letter
from openpyxl import Workbook

from usuarios.models import Usuario

from incidencias.models import (
    Incidencia,
    OrdenAtencion,
    Tecnico,
    Zona,
    CoordenadaZona,
    BoleteroCajero,
    Terminal,
    Cliente,
    TecnicoZona,
)


def _require_session(request):
    """Redirect to login si no hay sesión."""
    if not request.session.get('uid'):
        return redirect('login')
    return None

@require_http_methods(["GET", "POST"])
def login_view(request):
    # 🚫 SOLO redirigimos a home si es GET y ya hay sesión
    if request.method == "GET" and request.session.get('uid'):
        return redirect('home')

    if request.method == "POST":
        usuario = (request.POST.get('usuario') or '').strip()
        contrasena = (request.POST.get('contrasena') or '').strip()

        user = Usuario.objects.filter(
            usuario_login__iexact=usuario,
            contrasena=contrasena
        ).first()

        if not user:
            messages.error(request, "Usuario o contraseña incorrectos.")
        else:
            # Opcional pero recomendable: limpiar sesión anterior
            request.session.flush()

            request.session['uid']      = user.id_usuario
            request.session['nombre']   = user.nombre
            request.session['rol']      = user.rol
            request.session['terminal'] = user.id_terminal_id  # puede ser None

            return redirect('home')

    # Si es GET sin sesión, o POST con error → mostrar login
    return render(request, 'login_full.html')


def logout_view(request):
    request.session.flush()
    return redirect('login')


def dashboard_redirect(request):
    if not request.session.get('uid'):
        return redirect('login')

    rol = (request.session.get('rol') or '').lower()

    if rol == 'control_interno':
        return redirect('panel_control_interno')

    elif rol == 'tecnico':
        return redirect('panel_tecnico')

    elif rol in ('admin_terminal', 'terminal'):
        return redirect('panel_admin_terminal')

    elif rol in ('admin_sistema', 'admin'):
        return redirect('panel_admin_sistema')

    return redirect('panel_control_interno')



# ============================
#  PANEL CONTROL INTERNO (UI)
# ============================
# pega esto reemplazando la función panel_control_interno existente
def panel_control_interno(request):
    maybe_redirect = _require_session(request)
    if maybe_redirect:
        return maybe_redirect

    f_asig_desde = request.GET.get("f_asig_desde", "")
    f_asig_hasta = request.GET.get("f_asig_hasta", "")
    f_aten_desde = request.GET.get("f_aten_desde", "")
    f_aten_hasta = request.GET.get("f_aten_hasta", "")
    estado = request.GET.get("estado", "")
    tecnico_selected = request.GET.get("tecnico", "")
    zona_selected = request.GET.get("zona", "")
    codigo_cliente = request.GET.get("codigo_cliente", "")

    qs = (
        OrdenAtencion.objects
        .select_related(
            "id_cliente",
            "id_tecnico",
            "id_tecnico__id_usuario",
            "id_zona"
        )
        .order_by("-fecha_asignacion", "-id_orden")
    )

    if f_asig_desde:
        qs = qs.filter(fecha_asignacion__date__gte=f_asig_desde)

    if f_asig_hasta:
        qs = qs.filter(fecha_asignacion__date__lte=f_asig_hasta)

    if f_aten_desde:
        qs = qs.filter(fecha_atencion__date__gte=f_aten_desde)

    if f_aten_hasta:
        qs = qs.filter(fecha_atencion__date__lte=f_aten_hasta)

    if estado:
        qs = qs.filter(estado=estado)

    if tecnico_selected:
        qs = qs.filter(id_tecnico_id=tecnico_selected)

    if zona_selected:
        qs = qs.filter(id_zona_id=zona_selected)

    if codigo_cliente:
        qs = qs.filter(id_cliente__codigo_cliente__icontains=codigo_cliente)

    qs = qs[:100]

    paginator = Paginator(qs, 20)
    page_num = request.GET.get("page", 1)

    try:
        page_obj = paginator.page(page_num)
    except:
        page_obj = paginator.page(1)

    rows = []

    for orden in page_obj.object_list:
        cliente = orden.id_cliente
        tecnico = orden.id_tecnico

        rows.append({
            "id_orden": orden.id_orden,
            "fecha_asignacion": orden.fecha_asignacion,
            "codigo_cliente": cliente.codigo_cliente if cliente else "—",
            "cliente": cliente.nombre_cliente if cliente else "—",
            "celular": cliente.celular if cliente else "—",
            "direccion": cliente.direccion if cliente else "—",
            "distrito": cliente.distrito if cliente else "—",
            "zona": orden.id_zona.nombre_zona if orden.id_zona else "Sin zona",

            "tecnico": (
                tecnico.id_usuario.nombre
                if tecnico and tecnico.id_usuario
                else "Sin asignar"
            ),

            "usuario_tecnico": (
                tecnico.id_usuario.usuario_login
                if tecnico and tecnico.id_usuario
                else "—"
            ),

            "estado": orden.estado,
            "indicaciones": orden.indicaciones or "—",
            "observacion_tecnico": orden.observacion_tecnico or "—",
            "fecha_atencion": orden.fecha_atencion,
        })

    preserved = request.GET.urlencode()
    if preserved:
        preserved = "&" + preserved

    context = {
        "usuario_nombre": request.session.get("nombre", "Usuario"),
        "terminal_actual": request.session.get("terminal") or "—",

        "rows": rows,
        "page_obj": page_obj,
        "preserved": preserved,

        "f_asig_desde": f_asig_desde,
        "f_asig_hasta": f_asig_hasta,
        "f_aten_desde": f_aten_desde,
        "f_aten_hasta": f_aten_hasta,
        "estado": estado,
        "tecnico_selected": tecnico_selected,
        "zona_selected": zona_selected,
        "codigo_cliente": codigo_cliente,

        "zonas": Zona.objects.filter(activo=True).order_by("nombre_zona"),
        "tecnicos": Tecnico.objects.select_related("id_usuario").order_by("id_usuario__nombre"),
    }

    return render(request, "control_interno_dashboard.html", context)

# ============================
#  PLACEHOLDERS (puedes cambiar luego)
# ============================
#def panel_admin_terminal(request):
 #   maybe_redirect = _require_session(request)
 #   if maybe_redirect: 
 #       return maybe_redirect
 #   return render(request, 'placeholder.html', {'titulo': 'Panel Admin Terminal'})

def panel_admin_terminal(request):
    # 1) Verificar sesión
    maybe_redirect = _require_session(request)
    if maybe_redirect:
        return maybe_redirect

    # 2) Obtener usuario logueado y su terminal
    uid = request.session.get('uid')
    user = (
        Usuario.objects
        .select_related('id_terminal')
        .filter(pk=uid)
        .first()
    )

    terminal_id = None
    terminal_name = '—'

    if user and user.id_terminal_id:
        terminal_id = user.id_terminal_id
        terminal_name = getattr(user.id_terminal, 'nombre_terminal', str(user.id_terminal_id))

    # ==========================================================
    # 3) Filtros (igual que Control Interno)
    #    Fecha de incidencia: f1, f2
    #    Fecha de revisión:  r1, r2
    # ==========================================================
    f1 = (request.GET.get('f1') or '').strip()  # YYYY-MM-DD
    f2 = (request.GET.get('f2') or '').strip()
    r1 = (request.GET.get('r1') or '').strip()
    r2 = (request.GET.get('r2') or '').strip()

    # 4) Query base: incidencias SOLO del terminal del usuario
    qs = (
        Incidencia.objects
        .select_related('id_bc', 'id_usuario', 'id_bc__id_terminal')
        .order_by('-fecha_revision', '-id_incidencia')
    )

    if terminal_id is not None:
        qs = qs.filter(id_bc__id_terminal_id=terminal_id)

    # =========================
    # 5) Aplicar filtros fechas
    # =========================
    # ---- Fecha de incidencia ----
    try:
        if f1:
            d1 = date.fromisoformat(f1)
            # si fecha_incidencia es DateField:
            qs = qs.filter(fecha_incidencia__gte=d1)
        if f2:
            d2 = date.fromisoformat(f2)
            qs = qs.filter(fecha_incidencia__lte=d2)
    except ValueError:
        pass

    # ---- Fecha de revisión ---- (normalmente DateTimeField)
    try:
        if r1:
            rd1 = date.fromisoformat(r1)
            qs = qs.filter(fecha_revision__gte=rd1)
        if r2:
            rd2 = date.fromisoformat(r2)
            qs = qs.filter(fecha_revision__lte=rd2)
    except ValueError:
        pass

    # 👉 Query para tarjetas (ya incluye filtros de fecha y terminal)
    qs_for_cards = qs

    # 6) Filtro por estado vía cards
    estado_card = (request.GET.get('estado_card') or '').lower()
    if estado_card in ('conforme', 'observado', 'pendiente', 'resuelto'):
        qs = qs.filter(estado__iexact=estado_card)

    # 7) Paginación
    per_page = 20
    paginator = Paginator(qs, per_page)

    try:
        page_num = int(request.GET.get('page', 1))
    except ValueError:
        page_num = 1

    try:
        page_obj = paginator.page(page_num)
    except EmptyPage:
        page_obj = paginator.page(1)

    # 8) Construcción de filas para la tabla
    rows = []
    for inc in page_obj.object_list:
        inc_id = getattr(inc, 'id_incidencia', None)

        # Fecha de incidencia (usa el campo real)
        fecha_incidencia = getattr(inc, 'fecha_incidencia', None)

        bc = getattr(inc, 'id_bc', None)
        bc_nombre = getattr(bc, 'nombre', '—') if bc else '—'
        bc_usuario = getattr(bc, 'usuario', '—') if bc else '—'
        bc_cargo = getattr(bc, 'cargo', '—') if bc else '—'

        term_obj = getattr(bc, 'id_terminal', None) if bc else None
        terminal_txt = getattr(term_obj, 'nombre_terminal', None) if term_obj else None
        if not terminal_txt:
            terminal_txt = '—'

        ci_obj = getattr(inc, 'id_usuario', None)
        control_interno = getattr(ci_obj, 'nombre', '—') if ci_obj else '—'

        estado_raw = (getattr(inc, 'estado', '') or '').lower()
        if estado_raw == 'observado':
            estado = 'Observado'
        elif estado_raw == 'pendiente':
            estado = 'Pendiente'
        elif estado_raw == 'resuelto':
            estado = 'Resuelto'
        else:
            estado = 'Conforme'

        motivo = getattr(inc, 'motivo', '—')

        evidencia_val = getattr(inc, 'evidencia', None)
        evidencia = evidencia_val.url if evidencia_val else ''

        fecha_revision = getattr(inc, 'fecha_revision', None)

        rows.append({
            'id_incidencia': inc_id,
            'fecha_incidencia': fecha_incidencia,
            'nombre': bc_nombre,
            'usuario': bc_usuario,
            'cargo': bc_cargo,
            'terminal': terminal_txt,
            'control_interno': control_interno,
            'estado': estado,
            'estado_raw': estado_raw,
            'motivo': motivo,
            'evidencia': evidencia,
            'fecha_revision': fecha_revision,
        })

    # 9) Tarjetas
    cards = {
        'activas': qs_for_cards.exclude(estado__iexact='resuelto').count(),
        'pendientes': qs_for_cards.filter(estado__iexact='pendiente').count(),
        'resueltas': qs_for_cards.filter(estado__iexact='resuelto').count(),
        'conformes': qs_for_cards.filter(estado__iexact='conforme').count(),
        'observadas': qs_for_cards.filter(estado__iexact='observado').count(),
        'total': qs_for_cards.count(),
    }

    # 10) Preservar filtros en paginación
    preserved = ''
    if f1:
        preserved += f'&f1={f1}'
    if f2:
        preserved += f'&f2={f2}'
    if r1:
        preserved += f'&r1={r1}'
    if r2:
        preserved += f'&r2={r2}'
    if estado_card:
        preserved += f'&estado_card={estado_card}'

    # 11) Contexto
    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'terminal_name': terminal_name,

        'rows': rows,
        'cards': cards,

        'page_obj': page_obj,
        'preserved': preserved,

        # valores para mantener inputs rellenos
        'f1': f1,
        'f2': f2,
        'r1': r1,
        'r2': r2,

        'estado_card': estado_card,
    }

    return render(request, 'admin_terminal_dashboard.html', context)





def panel_admin_sistema(request):
    if not request.session.get('uid'):
        return redirect('login')

    if request.method == 'POST':
        accion = (request.POST.get('accion') or 'create').lower()
        id_usuario = request.POST.get('id_usuario')

        nombre = (request.POST.get('nombre') or '').strip()
        usuario_login = (request.POST.get('usuario_login') or '').strip()
        contrasena = (request.POST.get('contrasena') or '').strip()
        rol = (request.POST.get('rol') or '').strip()
        activo = bool(request.POST.get('activo'))
        terminal_id = request.POST.get('id_terminal') or None
        celular = (request.POST.get('celular') or '').strip()
        turno = (request.POST.get('turno') or '').strip()

        if accion == 'delete' and id_usuario:
            usuario = Usuario.objects.filter(pk=id_usuario).first()
            if usuario:
                Tecnico.objects.filter(id_usuario=usuario).delete()
                usuario.delete()
            return redirect('panel_admin_sistema')

        if nombre and usuario_login and rol:

            if accion == 'create' and Usuario.objects.filter(usuario_login=usuario_login).exists():
                return redirect('panel_admin_sistema')

            if accion == 'update' and id_usuario:
                u = Usuario.objects.filter(pk=id_usuario).first()
                if not u:
                    return redirect('panel_admin_sistema')
            else:
                u = Usuario()
                if hasattr(u, 'fecha_ingreso'):
                    u.fecha_ingreso = timezone.now().date()

            u.nombre = nombre
            u.usuario_login = usuario_login

            if contrasena:
                u.contrasena = contrasena

            u.rol = rol

            if hasattr(u, 'activo'):
                u.activo = activo

            u.id_terminal_id = terminal_id if terminal_id else None
            u.save()

            if rol == "tecnico":
                tecnico, creado = Tecnico.objects.get_or_create(
                    id_usuario=u,
                    defaults={
                        "celular": celular,
                        "turno": turno,
                        "fecha_creacion": timezone.now()
                    }
                )

                tecnico.celular = celular
                tecnico.turno = turno
                tecnico.save()

            else:
                Tecnico.objects.filter(id_usuario=u).delete()

        return redirect('panel_admin_sistema')

    qs = Usuario.objects.select_related('id_terminal').order_by('id_usuario')

    paginator = Paginator(qs, 12)
    page_num = request.GET.get('page', 1)

    try:
        page_obj = paginator.page(page_num)
    except:
        page_obj = paginator.page(1)

    usuarios_rows = []

    roles = {
        "admin_sistema": "Administrador del Sistema",
        "coordinador": "Coordinador",
        "tecnico": "Técnico",
        "control_interno": "Control Interno",
        "admin_terminal": "Administrador de Terminal",
    }

    for u in page_obj.object_list:
        tecnico = Tecnico.objects.filter(id_usuario=u).first()

        activo_val = getattr(u, 'activo', True)
        activo_bool = bool(activo_val)

        usuarios_rows.append({
            'id': u.id_usuario,
            'fecha_ingreso': getattr(u, 'fecha_ingreso', None),
            'nombre': u.nombre,
            'usuario_login': u.usuario_login,
            'contrasena': u.contrasena or '',
            'celular': tecnico.celular if tecnico else '',
            'turno': tecnico.turno if tecnico else '',
            'rol': u.rol,
            'rol_mostrar': roles.get(u.rol, u.rol),
            'terminal_id': getattr(u, 'id_terminal_id', '') or '',
            'activo': activo_bool,
            'activo_raw': '1' if activo_bool else '0',
        })

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'usuarios': usuarios_rows,
        'page_obj': page_obj,
        'terminales': Terminal.objects.all().order_by('nombre_terminal'),
        'hoy': timezone.now().date(),
    }

    return render(request, 'admin_sistema_dashboard.html', context)


#def panel_admin_sistema(request):
 #   maybe_redirect = _require_session(request)
 #   if maybe_redirect: 
 #       return maybe_redirect
 #   return render(request, 'placeholder.html', {'titulo': 'Panel Admin Sistema'})










def exportar_incidencias_excel(request):
    """
    Exporta las órdenes de atención según los filtros actuales del dashboard.
    """

    f_asig_desde = request.GET.get("f_asig_desde", "")
    f_asig_hasta = request.GET.get("f_asig_hasta", "")
    f_aten_desde = request.GET.get("f_aten_desde", "")
    f_aten_hasta = request.GET.get("f_aten_hasta", "")
    estado = request.GET.get("estado", "")
    tecnico_selected = request.GET.get("tecnico", "")
    zona_selected = request.GET.get("zona", "")
    codigo_cliente = request.GET.get("codigo_cliente", "")

    qs = (
        OrdenAtencion.objects
        .select_related(
            "id_cliente",
            "id_tecnico",
            "id_tecnico__id_usuario",
            "id_zona"
        )
        .order_by("-fecha_asignacion", "-id_orden")
    )

    if f_asig_desde:
        qs = qs.filter(fecha_asignacion__date__gte=f_asig_desde)

    if f_asig_hasta:
        qs = qs.filter(fecha_asignacion__date__lte=f_asig_hasta)

    if f_aten_desde:
        qs = qs.filter(fecha_atencion__date__gte=f_aten_desde)

    if f_aten_hasta:
        qs = qs.filter(fecha_atencion__date__lte=f_aten_hasta)

    if estado:
        qs = qs.filter(estado=estado)

    if tecnico_selected:
        qs = qs.filter(id_tecnico_id=tecnico_selected)

    if zona_selected:
        qs = qs.filter(id_zona_id=zona_selected)

    if codigo_cliente:
        qs = qs.filter(id_cliente__codigo_cliente__icontains=codigo_cliente)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ordenes de atencion"

    headers = [
        "Fecha asignación",
        "Hora asignación",
        "Código cliente",
        "Cliente",
        "Celular",
        "Dirección",
        "Distrito",
        "Zona",
        "Técnico",
        "Usuario técnico",
        "Estado",
        "Indicaciones",
        "Fecha atención",
        "Hora atención",
        "Observación técnico",
    ]

    ws.append(headers)

    for orden in qs:
        cliente = orden.id_cliente
        tecnico = orden.id_tecnico

        fecha_asig = orden.fecha_asignacion
        fecha_aten = orden.fecha_atencion

        ws.append([
            fecha_asig.strftime("%d-%m-%Y") if fecha_asig else "",
            fecha_asig.strftime("%H:%M") if fecha_asig else "",
            cliente.codigo_cliente if cliente else "",
            cliente.nombre_cliente if cliente else "",
            cliente.celular if cliente else "",
            cliente.direccion if cliente else "",
            cliente.distrito if cliente else "",
            orden.id_zona.nombre_zona if orden.id_zona else "Sin zona",
            tecnico.id_usuario.nombre if tecnico and tecnico.id_usuario else "Sin asignar",
            tecnico.id_usuario.usuario_login if tecnico and tecnico.id_usuario else "",
            "Por atender" if orden.estado == "por_atender" else "Atendido" if orden.estado == "atendido" else orden.estado,
            orden.indicaciones or "",
            fecha_aten.strftime("%d-%m-%Y") if fecha_aten else "",
            fecha_aten.strftime("%H:%M") if fecha_aten else "",
            orden.observacion_tecnico or "",
        ])

    column_widths = [18, 15, 18, 30, 15, 40, 20, 25, 30, 20, 15, 35, 18, 15, 35]

    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "ordenes_atencion_export.xlsx"

    response = HttpResponse(
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response





@require_POST
def resolver_incidencia(request):
    # Verificar sesión como en el panel
    maybe_redirect = _require_session(request)
    if maybe_redirect:
        return maybe_redirect

    inc_id = request.POST.get('id_incidencia')
    incidencia = get_object_or_404(Incidencia, pk=inc_id)

    # Campos del formulario
    solucion_txt = request.POST.get('solucion_admin', '').strip()
    fecha_sol_str = request.POST.get('fecha_solucion', '')
    evidencia_file = request.FILES.get('evidencia_solucion')

    # Guardar solución
    incidencia.solucion_admin = solucion_txt

    # ⭐ FECHA DE SOLUCIÓN → FECHA LOCAL DE PERÚ (CORRECCIÓN)
    incidencia.fecha_solucion = timezone.localdate()

    # Evidencia de solución (si adjuntan)
    if evidencia_file:
        incidencia.evidencia_solucion = evidencia_file

    # Actualizar estado
    incidencia.estado = 'pendiente'

    incidencia.save()

    messages.success(request, 'La incidencia fue actualizada correctamente.')
    return redirect('panel_admin_terminal')






@require_POST
def actualizar_estado_control_interno(request):
    inc_id = request.POST.get('id_incidencia')
    nuevo_estado = request.POST.get('nuevo_estado')

    incidencia = get_object_or_404(Incidencia, pk=inc_id)

    # Control interno valida los estados finales
    if nuevo_estado == "resuelto":
        incidencia.estado = "resuelto"
    elif nuevo_estado == "observado":
        incidencia.estado = "observado"

    
    incidencia.save()

    messages.success(request, "Estado actualizado correctamente.")
    return redirect("panel_control_interno")




def panel_incidencias(request):
    maybe_redirect = _require_session(request)
    if maybe_redirect:
        return maybe_redirect

    # ==========================
    # 🔍 Filtros GET
    # ==========================
    produccion_hoy = request.GET.get("produccion_hoy")
    # Fecha incidencia
    f1 = request.GET.get("f1")   # desde
    f2 = request.GET.get("f2")   # hasta

    # Fecha revisión
    r1 = request.GET.get("r1")
    r2 = request.GET.get("r2")

    terminal_filter = request.GET.get("terminal")
    estado_filter = request.GET.get("estado")
    usuario_filter = (request.GET.get("usuario") or "").strip()
    motivo_filter = (request.GET.get("motivo") or "").strip()
    ci_filter = (request.GET.get("ci") or "").strip()

    # ==========================
    # Base Query
    # ==========================

    qs = (
        Incidencia.objects
        .select_related("id_bc", "id_usuario", "id_bc__id_terminal")
        .order_by("-fecha_revision", "-id_incidencia")
    )

    # ==========================
    # 🎯 Filtro FECHA INCIDENCIA
    # ==========================
    if f1 and f2:
        try:
            d1 = date.fromisoformat(f1)
            d2 = date.fromisoformat(f2)
            if d1 > d2:
                d1, d2 = d2, d1
            qs = qs.filter(fecha_incidencia__range=(d1, d2))
        except:
            pass

    # ==========================
    # 🎯 Filtro FECHA REVISIÓN
    # ==========================
    if r1 and r2:
        try:
            d1 = date.fromisoformat(r1)
            d2 = date.fromisoformat(r2)
            if d1 > d2:
                d1, d2 = d2, d1
            qs = qs.filter(fecha_revision__range=(d1, d2))
        except:
            pass

    # ==========================
    # 🎯 Filtro Terminal
    # ==========================
    if terminal_filter:
        qs = qs.filter(id_bc__id_terminal_id=terminal_filter)

    # ==========================
    # 🎯 Filtro Estado NORMALIZADO
    # ==========================
    if estado_filter:
        # viene como "Conforme", "Observado", "Pendiente", "Resuelto"
        estado_filter_norm = (estado_filter or "").lower().replace("ó", "o")

        if estado_filter_norm.startswith("observ"):
            qs = qs.filter(estado__icontains="observ")
        elif estado_filter_norm.startswith("pend"):
            qs = qs.filter(estado__icontains="pendiente")
        elif estado_filter_norm.startswith("resu"):
            qs = qs.filter(estado__icontains="resuelto")
        elif estado_filter_norm.startswith("conf"):
            qs = qs.filter(estado__icontains="conforme")


    # ==========================
    # 🎯 Usuario boletero/cajero
    # ==========================
    if usuario_filter:
        qs = qs.filter(
            Q(id_bc__nombre__icontains=usuario_filter) |
            Q(id_bc__usuario__icontains=usuario_filter)
        )

 # 🎯 FILTRO: Producción hoy
    if produccion_hoy:
        today = timezone.localdate()
        qs = qs.filter(fecha_revision=today)

    # ==========================
    # 🎯 Control interno
    # ==========================
    if ci_filter:
        qs = qs.filter(id_usuario__usuario_login__icontains=ci_filter)


    # ==========================
    #  LIMIT TOP 100
    # ==========================
    qs = qs[:100]

    # ==========================
    #  Paginación
    # ==========================
    paginator = Paginator(qs, 20)
    page_num = request.GET.get("page", 1)

    try:
        page_obj = paginator.page(page_num)
    except:
        page_obj = paginator.page(1)

    # ==========================
    #  Construcción de tabla rows
    # ==========================
    rows = []

    for inc in page_obj.object_list:

        bc = inc.id_bc
        term = bc.id_terminal if bc else None

        # NORMALIZAR ESTADO (4 estados posibles)
        estado_val = (inc.estado or "").lower().replace("ó", "o")

        if estado_val.startswith("observ"):
            estado = "Observado"
        elif estado_val.startswith("pend"):
            estado = "Pendiente"
        elif estado_val.startswith("resu"):
            estado = "Resuelto"
        else:
            estado = "Conforme"


        evidencia = inc.evidencia.url if inc.evidencia else ""

        rows.append({
            "id_incidencia": inc.id_incidencia,  # 👈 NECESARIO PARA EL POST
            "fecha": inc.fecha_incidencia,
            "nombre": bc.nombre if bc else "—",
            "usuario": bc.usuario if bc else "—",
            "cargo": bc.cargo if bc else "—",
            "terminal": term.nombre_terminal if term else "—",
            "control_interno": inc.id_usuario.nombre if inc.id_usuario else "—",
            "estado": estado,
            "motivo": inc.motivo or "—",
            "evidencia": evidencia,
            "fecha_revision": inc.fecha_revision,
            "solucion_admin": inc.solucion_admin or "",  # 👈 para el Ver respuesta
            "evidencia_solucion": (
                inc.evidencia_solucion.url if getattr(inc, "evidencia_solucion", None) else ""
            ),  # 👈 evidencia de solución
        })

    # ==========================
    #  Preservar filtros para paginación
    # ==========================
    preserved = ""
    for k in ("f1", "f2", "r1", "r2", "terminal", "estado", "usuario", "motivo", "ci"):
        v = request.GET.get(k)
        if v:
            preserved += f"&{k}={v}"

    # ==========================
    # Contexto final
    # ==========================
    context = {
        "usuario_nombre": request.session.get("nombre", "Usuario"),
        "terminal_actual": request.session.get("terminal") or "—",

        "rows": rows,
        "page_obj": page_obj,
        "preserved": preserved,

        # reutilizamos valores para mantenerlos en los inputs
        "f1": f1 or "",
        "f2": f2 or "",
        "r1": r1 or "",
        "r2": r2 or "",
        "terminal_selected": terminal_filter or "",
        "estado": estado_filter or "",
        "usuario": usuario_filter or "",
        "motivo": motivo_filter or "",
        "ci": ci_filter or "",

        "boleteros": BoleteroCajero.objects.all(),
        "terminales": Terminal.objects.all().order_by("id_terminal"),

        # NUEVOS SELECTS
        "usuarios_bc": BoleteroCajero.objects.filter(estado="activo").order_by("usuario"),
        "usuarios_ci": Usuario.objects.filter(rol="control_interno").order_by("nombre"),
            
        "boleteros": BoleteroCajero.objects.all(),
        "terminales": Terminal.objects.all().order_by("id_terminal"),
        "control_internos": Usuario.objects.filter(rol="control_interno"),
        
        
    }
    return render(request, 'admin_incidencias_dashboard.html', context)









def panel_boleteros(request):
    if not request.session.get('uid'):
        return redirect('login')

    qs = (
        Tecnico.objects
        .select_related('id_usuario')
        .order_by('id_usuario__nombre')
    )

    paginator = Paginator(qs, 15)
    page_num = request.GET.get('page', 1)

    try:
        page_obj = paginator.page(page_num)
    except:
        page_obj = paginator.page(1)

    tecnicos_rows = []

    for t in page_obj.object_list:
        usuario = t.id_usuario

        asignaciones = (
            TecnicoZona.objects
            .select_related('id_zona')
            .filter(
                id_tecnico=t,
                activo=True
            )
        )

        zonas_asignadas = [
            a.id_zona
            for a in asignaciones
            if a.id_zona
        ]

        zona_principal = zonas_asignadas[0] if zonas_asignadas else None

        tecnicos_rows.append({
            'id_tecnico': t.id_tecnico,
            'fecha_creacion': t.fecha_creacion,
            'nombre': usuario.nombre if usuario else '—',
            'usuario_login': usuario.usuario_login if usuario else '—',
            'celular': t.celular or '—',
            'turno': t.turno or 'Sin turno',
            'estado': 'Activo' if usuario and usuario.activo else 'Inactivo',
            'latitud': t.latitud_actual or '—',
            'longitud': t.longitud_actual or '—',
            'tiene_zona': True if zonas_asignadas else False,
            'nombre_zona': ', '.join([z.nombre_zona for z in zonas_asignadas]) if zonas_asignadas else 'Sin zona',
            'id_zona': zona_principal.id_zona if zona_principal else None,
        })

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'tecnicos_rows': tecnicos_rows,
        'page_obj': page_obj,
        'zonas': Zona.objects.filter(activo=True).order_by('nombre_zona'),
    }

    return render(
        request,
        'admin_boleteros_dashboard.html',
        context
    )



def asignar_zona_tecnico(request, id_tecnico):
    if not request.session.get('uid'):
        return redirect('login')

    tecnico = get_object_or_404(Tecnico, pk=id_tecnico)

    zona = Zona.objects.filter(id_tecnico=tecnico).first()

    coordenadas_zona = []

    if zona:
        puntos = CoordenadaZona.objects.filter(
            id_zona=zona
        ).order_by('orden_punto')

        coordenadas_zona = [
            [float(p.latitud), float(p.longitud)]
            for p in puntos
        ]

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'tecnico': tecnico,
        'zona': zona,
        'coordenadas_zona': coordenadas_zona,
    }

    return render(
        request,
        'asignar_zona_tecnico.html',
        context
    )





@require_POST
def guardar_zona_tecnico(request, id_tecnico):
    if not request.session.get('uid'):
        return JsonResponse({'ok': False, 'error': 'Sesión no válida'}, status=403)

    tecnico = get_object_or_404(Tecnico, pk=id_tecnico)

    try:
        data = json.loads(request.body)
        nombre_zona = (data.get('nombre_zona') or '').strip()
        coordenadas = data.get('coordenadas') or []

        if not nombre_zona:
            return JsonResponse({'ok': False, 'error': 'Ingrese el nombre de la zona'})

        if len(coordenadas) < 3:
            return JsonResponse({'ok': False, 'error': 'Debe dibujar al menos 3 puntos'})

        zona = Zona.objects.filter(id_tecnico=tecnico).first()

        if zona:
            zona.nombre_zona = nombre_zona
            zona.fecha_creacion = timezone.now()
            zona.save()

            CoordenadaZona.objects.filter(id_zona=zona).delete()
        else:
            zona = Zona.objects.create(
                nombre_zona=nombre_zona,
                id_tecnico=tecnico,
                fecha_creacion=timezone.now()
            )

        for punto in coordenadas:
            CoordenadaZona.objects.create(
                id_zona=zona,
                latitud=punto.get('latitud'),
                longitud=punto.get('longitud'),
                orden_punto=punto.get('orden')
            )

        return JsonResponse({'ok': True})

    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=500)
    
    
    
    


def ver_zona_tecnico(request, id_tecnico):
    if not request.session.get('uid'):
        return redirect('login')

    tecnico = get_object_or_404(Tecnico, pk=id_tecnico)

    zona = Zona.objects.filter(id_tecnico=tecnico).first()
    coordenadas_zona = []

    if zona:
        puntos = CoordenadaZona.objects.filter(id_zona=zona).order_by('orden_punto')
        coordenadas_zona = [
            [float(p.latitud), float(p.longitud)]
            for p in puntos
        ]

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'tecnico': tecnico,
        'zona': zona,
        'coordenadas_zona': coordenadas_zona,
    }

    return render(request, 'ver_zona_tecnico.html', context)



def buscar_cliente(request):
    codigo = request.GET.get("codigo_cliente", "").strip()

    cliente = Cliente.objects.filter(codigo_cliente=codigo).first()

    if not cliente:
        return JsonResponse({"existe": False})

    return JsonResponse({
        "existe": True,
        "nombre_cliente": cliente.nombre_cliente or "",
        "celular": cliente.celular or "",
        "direccion": cliente.direccion or "",
        "distrito": cliente.distrito or "",
        "departamento": cliente.departamento or "",
        "latitud": str(cliente.latitud) if cliente.latitud else "",
        "longitud": str(cliente.longitud) if cliente.longitud else "",
    })
    
    
    
def panel_tecnico(request):
    if not request.session.get('uid'):
        return redirect('login')

    uid = request.session.get('uid')

    usuario = Usuario.objects.filter(pk=uid).first()
    tecnico = Tecnico.objects.filter(id_usuario=usuario).first()

    if not tecnico:
        return redirect('login')

    qs = (
        OrdenAtencion.objects
        .select_related("id_cliente", "id_tecnico", "id_tecnico__id_usuario")
        .filter(id_tecnico=tecnico)
        .order_by("-fecha_asignacion", "-id_orden")
    )

    rows = []

    for orden in qs:
        cliente = orden.id_cliente

        rows.append({
            "id_orden": orden.id_orden,
            "fecha_asignacion": orden.fecha_asignacion,
            "codigo_cliente": cliente.codigo_cliente,
            "cliente": cliente.nombre_cliente,
            "celular": cliente.celular,
            "direccion": cliente.direccion,
            "distrito": cliente.distrito,
            "departamento": cliente.departamento,
            "latitud": cliente.latitud,
            "longitud": cliente.longitud,
            "tecnico": tecnico.id_usuario.nombre,
            "estado": orden.estado,
            "indicaciones": orden.indicaciones,
            "fecha_atencion": orden.fecha_atencion,
            "observacion_tecnico": orden.observacion_tecnico,
        })

    context = {
        "usuario_nombre": usuario.nombre,
        "rows": rows,
    }

    return render(request, "tecnico_dashboard.html", context)




@require_POST
def atender_orden_tecnico(request):
    id_orden = request.POST.get("id_orden")
    observacion = request.POST.get("observacion_tecnico", "").strip()

    fecha_manual = request.POST.get("fecha_atencion_manual", "").strip()
    hora_manual = request.POST.get("hora_atencion_manual", "").strip()

    orden = OrdenAtencion.objects.get(pk=id_orden)

    fecha_atencion_final = datetime.now()

    if fecha_manual and hora_manual:
        fecha_hora_str = f"{fecha_manual} {hora_manual}"
        fecha_atencion_final = datetime.strptime(
            fecha_hora_str,
            "%Y-%m-%d %H:%M"
        )

    orden.estado = "atendido"
    orden.observacion_tecnico = observacion
    orden.fecha_atencion = fecha_atencion_final
    orden.save()

    return redirect("panel_tecnico")



@require_POST
def guardar_asignacion_zona(request):
    if not request.session.get('uid'):
        return redirect('login')

    id_tecnico = request.POST.get('id_tecnico')
    zonas_ids = request.POST.getlist('zonas')

    tecnico = get_object_or_404(Tecnico, pk=id_tecnico)

    TecnicoZona.objects.filter(
        id_tecnico=tecnico,
        activo=True
    ).update(activo=False)

    for id_zona in zonas_ids:
        zona = get_object_or_404(Zona, pk=id_zona)

        TecnicoZona.objects.create(
            id_tecnico=tecnico,
            id_zona=zona,
            fecha_asignacion=timezone.now(),
            activo=True
        )

    return redirect('panel_boleteros')




def panel_zonas(request):
    if not request.session.get('uid'):
        return redirect('login')

    zonas = Zona.objects.all().order_by('nombre_zona')

    zonas_rows = []

    for z in zonas:
        asignaciones = (
            TecnicoZona.objects
            .select_related('id_tecnico', 'id_tecnico__id_usuario')
            .filter(id_zona=z, activo=True)
        )

        tecnicos = []
        for a in asignaciones:
            if a.id_tecnico and a.id_tecnico.id_usuario:
                tecnicos.append(a.id_tecnico.id_usuario.nombre)

        zonas_rows.append({
            'id_zona': z.id_zona,
            'nombre_zona': z.nombre_zona,
            'fecha_creacion': z.fecha_creacion,
            'activo': z.activo,
            'tecnicos': ', '.join(tecnicos) if tecnicos else 'Sin técnicos',
        })

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'zonas_rows': zonas_rows,
    }

    return render(request, 'admin_zonas_dashboard.html', context)





def crear_zona(request):

    if not request.session.get('uid'):
        return redirect('login')

    return render(
        request,
        'crear_zona.html'
    )




@require_POST
def guardar_zona(request):
    if not request.session.get('uid'):
        return JsonResponse({'ok': False, 'error': 'Sesión no válida'}, status=403)

    try:
        data = json.loads(request.body)

        nombre_zona = (data.get('nombre_zona') or '').strip()
        coordenadas = data.get('coordenadas') or []

        if not nombre_zona:
            return JsonResponse({'ok': False, 'error': 'Ingrese el nombre de la zona.'})

        if len(coordenadas) < 3:
            return JsonResponse({'ok': False, 'error': 'Debe dibujar una zona con al menos 3 puntos.'})

        zona = Zona.objects.create(
            nombre_zona=nombre_zona,
            id_tecnico=None,
            fecha_creacion=timezone.now(),
            activo=True
        )

        puntos = []

        for punto in coordenadas:
            puntos.append(
                CoordenadaZona(
                    id_zona=zona,
                    latitud=punto.get('latitud'),
                    longitud=punto.get('longitud'),
                    orden_punto=punto.get('orden')
                )
            )

        CoordenadaZona.objects.bulk_create(puntos)

        return JsonResponse({'ok': True, 'id_zona': zona.id_zona})

    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=500)
    
    
    
    




def ver_zona(request, id_zona):
    if not request.session.get('uid'):
        return redirect('login')

    zona = get_object_or_404(Zona, pk=id_zona)

    puntos = CoordenadaZona.objects.filter(
        id_zona=zona
    ).order_by('orden_punto')

    coordenadas_zona = [
        [float(p.latitud), float(p.longitud)]
        for p in puntos
    ]
    origen = request.GET.get("origen", "zonas")

    volver_url = "panel_boleteros" if origen == "tecnicos" else "panel_zonas"

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'zona': zona,
        'coordenadas_zona': coordenadas_zona,
        "volver_url": volver_url,
    }

    return render(request, 'ver_zona.html', context)


def editar_zona(request, id_zona):
    if not request.session.get('uid'):
        return redirect('login')

    zona = get_object_or_404(Zona, pk=id_zona)

    puntos = CoordenadaZona.objects.filter(
        id_zona=zona
    ).order_by('orden_punto')

    coordenadas_zona = [
        [float(p.latitud), float(p.longitud)]
        for p in puntos
    ]

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'zona': zona,
        'coordenadas_zona': coordenadas_zona,
    }

    return render(request, 'editar_zona.html', context)


@require_POST
def actualizar_zona(request, id_zona):
    if not request.session.get('uid'):
        return JsonResponse({'ok': False, 'error': 'Sesión no válida'}, status=403)

    zona = get_object_or_404(Zona, pk=id_zona)

    try:
        data = json.loads(request.body)

        nombre_zona = (data.get('nombre_zona') or '').strip()
        coordenadas = data.get('coordenadas') or []

        if not nombre_zona:
            return JsonResponse({'ok': False, 'error': 'Ingrese el nombre de la zona.'})

        if len(coordenadas) < 3:
            return JsonResponse({'ok': False, 'error': 'Debe dibujar una zona con al menos 3 puntos.'})

        zona.nombre_zona = nombre_zona
        zona.save()

        CoordenadaZona.objects.filter(id_zona=zona).delete()

        for punto in coordenadas:
            CoordenadaZona.objects.create(
                id_zona=zona,
                latitud=punto.get('latitud'),
                longitud=punto.get('longitud'),
                orden_punto=punto.get('orden')
            )

        return JsonResponse({'ok': True})

    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=500)
    
    
    
    



def ver_zonas_tecnico(request, id_tecnico):
    if not request.session.get('uid'):
        return redirect('login')

    tecnico = get_object_or_404(Tecnico, pk=id_tecnico)

    asignaciones = (
        TecnicoZona.objects
        .select_related('id_zona')
        .filter(id_tecnico=tecnico, activo=True)
    )

    zonas_data = []

    for asignacion in asignaciones:
        zona = asignacion.id_zona

        puntos = CoordenadaZona.objects.filter(
            id_zona=zona
        ).order_by('orden_punto')

        coordenadas = [
            [float(p.latitud), float(p.longitud)]
            for p in puntos
        ]

        zonas_data.append({
            'nombre': zona.nombre_zona,
            'coordenadas': coordenadas
        })

    context = {
        'usuario_nombre': request.session.get('nombre', 'Usuario'),
        'tecnico': tecnico,
        'zonas_data': zonas_data,
    }

    return render(request, 'ver_zonas_tecnico.html', context)









def panel_ordenes_admin(request):
    maybe_redirect = _require_session(request)
    if maybe_redirect:
        return maybe_redirect

    f_asig_desde = request.GET.get("f_asig_desde", "")
    f_asig_hasta = request.GET.get("f_asig_hasta", "")
    f_aten_desde = request.GET.get("f_aten_desde", "")
    f_aten_hasta = request.GET.get("f_aten_hasta", "")
    estado = request.GET.get("estado", "")
    tecnico_selected = request.GET.get("tecnico", "")
    zona_selected = request.GET.get("zona", "")
    codigo_cliente = request.GET.get("codigo_cliente", "")

    qs = (
        OrdenAtencion.objects
        .select_related(
            "id_cliente",
            "id_tecnico",
            "id_tecnico__id_usuario",
            "id_zona"
        )
        .order_by("-fecha_asignacion", "-id_orden")
    )

    if f_asig_desde:
        qs = qs.filter(fecha_asignacion__date__gte=f_asig_desde)

    if f_asig_hasta:
        qs = qs.filter(fecha_asignacion__date__lte=f_asig_hasta)

    if f_aten_desde:
        qs = qs.filter(fecha_atencion__date__gte=f_aten_desde)

    if f_aten_hasta:
        qs = qs.filter(fecha_atencion__date__lte=f_aten_hasta)

    if estado:
        qs = qs.filter(estado=estado)

    if tecnico_selected:
        qs = qs.filter(id_tecnico_id=tecnico_selected)

    if zona_selected:
        qs = qs.filter(id_zona_id=zona_selected)

    if codigo_cliente:
        qs = qs.filter(id_cliente__codigo_cliente__icontains=codigo_cliente)

    qs = qs[:100]

    paginator = Paginator(qs, 20)
    page_num = request.GET.get("page", 1)

    try:
        page_obj = paginator.page(page_num)
    except:
        page_obj = paginator.page(1)

    rows = []

    for orden in page_obj.object_list:
        cliente = orden.id_cliente
        tecnico = orden.id_tecnico

        rows.append({
            "id_orden": orden.id_orden,
            "fecha_asignacion": orden.fecha_asignacion,
            "codigo_cliente": cliente.codigo_cliente if cliente else "—",
            "cliente": cliente.nombre_cliente if cliente else "—",
            "celular": cliente.celular if cliente else "—",
            "direccion": cliente.direccion if cliente else "—",
            "distrito": cliente.distrito if cliente else "—",
            "zona": orden.id_zona.nombre_zona if orden.id_zona else "Sin zona",
            "tecnico": tecnico.id_usuario.nombre if tecnico and tecnico.id_usuario else "Sin asignar",
            "usuario_tecnico": tecnico.id_usuario.usuario_login if tecnico and tecnico.id_usuario else "—",
            "estado": orden.estado,
            "indicaciones": orden.indicaciones or "—",
            "observacion_tecnico": orden.observacion_tecnico or "—",
            "fecha_atencion": orden.fecha_atencion,
        })

    preserved = request.GET.urlencode()
    if preserved:
        preserved = "&" + preserved

    context = {
        "usuario_nombre": request.session.get("nombre", "Usuario"),
        "rows": rows,
        "page_obj": page_obj,
        "preserved": preserved,

        "f_asig_desde": f_asig_desde,
        "f_asig_hasta": f_asig_hasta,
        "f_aten_desde": f_aten_desde,
        "f_aten_hasta": f_aten_hasta,
        "estado": estado,
        "tecnico_selected": tecnico_selected,
        "zona_selected": zona_selected,
        "codigo_cliente": codigo_cliente,

        "zonas": Zona.objects.filter(activo=True).order_by("nombre_zona"),
        "tecnicos": Tecnico.objects.select_related("id_usuario").order_by("id_usuario__nombre"),
    }

    return render(request, "admin_ordenes_dashboard.html", context)





@require_POST
def eliminar_orden(request, id_orden):
    maybe_redirect = _require_session(request)
    if maybe_redirect:
        return maybe_redirect

    orden = get_object_or_404(OrdenAtencion, pk=id_orden)
    orden.delete()

    return redirect("panel_control_interno")






def obtener_turno_por_fecha(fecha_hora):
    hora = fecha_hora.hour

    # 22:00 hasta 13:59 -> turno mañana
    if hora >= 22 or hora < 14:
        return "mañana"

    # 14:00 hasta 21:59 -> turno tarde
    return "tarde"




def punto_en_poligono(lat, lon, poligono):
    dentro = False
    j = len(poligono) - 1

    for i in range(len(poligono)):
        lat_i, lon_i = poligono[i]
        lat_j, lon_j = poligono[j]

        if ((lon_i > lon) != (lon_j > lon)):
            interseccion = (lat_j - lat_i) * (lon - lon_i) / (lon_j - lon_i) + lat_i

            if lat < interseccion:
                dentro = not dentro

        j = i

    return dentro



def obtener_orden_editar(request, id_orden):
    orden = get_object_or_404(
        OrdenAtencion.objects.select_related("id_cliente", "id_tecnico", "id_zona"),
        pk=id_orden
    )

    cliente = orden.id_cliente

    return JsonResponse({
        "id_orden": orden.id_orden,
        "codigo_cliente": cliente.codigo_cliente,
        "nombre_cliente": cliente.nombre_cliente,
        "celular": cliente.celular or "",
        "direccion": cliente.direccion,
        "distrito": cliente.distrito or "",
        "departamento": cliente.departamento or "",
        "latitud": str(cliente.latitud or ""),
        "longitud": str(cliente.longitud or ""),
        "indicaciones": orden.indicaciones or "",
        "estado": orden.estado,
        "id_tecnico": orden.id_tecnico.id_tecnico if orden.id_tecnico else "",
        "fecha_asignacion": orden.fecha_asignacion.strftime("%Y-%m-%d") if orden.fecha_asignacion else "",
        "hora_asignacion": orden.fecha_asignacion.strftime("%H:%M") if orden.fecha_asignacion else "",
    })


@require_POST
def actualizar_orden(request, id_orden):
    orden = get_object_or_404(OrdenAtencion, pk=id_orden)
    cliente = orden.id_cliente

    codigo_cliente = request.POST.get("codigo_cliente", "").strip()
    nombre_cliente = request.POST.get("nombre_cliente", "").strip()
    celular = request.POST.get("celular", "").strip()
    direccion = request.POST.get("direccion", "").strip()
    distrito = request.POST.get("distrito", "").strip()
    departamento = request.POST.get("departamento", "").strip()
    indicaciones = request.POST.get("indicaciones", "").strip()
    estado = request.POST.get("estado", "por_atender")

    latitud = request.POST.get("latitud") or None
    longitud = request.POST.get("longitud") or None

    fecha_manual = request.POST.get("fecha_asignacion", "")
    hora_manual = request.POST.get("hora_asignacion", "")

    modo_asignacion = request.POST.get("modo_asignacion", "automatica")
    tecnico_manual = request.POST.get("tecnico_manual")

    fecha_asignacion_final = orden.fecha_asignacion or datetime.now()

    if fecha_manual and hora_manual:
        fecha_asignacion_final = datetime.strptime(
            f"{fecha_manual} {hora_manual}",
            "%Y-%m-%d %H:%M"
        )

    cliente.codigo_cliente = codigo_cliente
    cliente.nombre_cliente = nombre_cliente
    cliente.celular = celular
    cliente.direccion = direccion
    cliente.distrito = distrito
    cliente.departamento = departamento
    cliente.latitud = latitud
    cliente.longitud = longitud
    cliente.save()

    orden.indicaciones = indicaciones
    orden.estado = estado
    orden.fecha_asignacion = fecha_asignacion_final

    if modo_asignacion == "manual" and tecnico_manual:
        orden.id_tecnico = get_object_or_404(Tecnico, pk=tecnico_manual)

    else:
        tecnico_asignado = None
        zona_asignada = None

        if latitud and longitud:
            lat_cliente = float(latitud)
            lon_cliente = float(longitud)

            zonas_coincidentes = []

            for zona in Zona.objects.filter(activo=True):
                puntos = CoordenadaZona.objects.filter(id_zona=zona).order_by("orden_punto")

                poligono = [
                    (float(p.latitud), float(p.longitud))
                    for p in puntos
                ]

                if len(poligono) >= 3 and punto_en_poligono(lat_cliente, lon_cliente, poligono):
                    zonas_coincidentes.append(zona)

            turno_actual = obtener_turno_por_fecha(fecha_asignacion_final)

            asignaciones = TecnicoZona.objects.select_related(
                "id_tecnico", "id_zona"
            ).filter(
                id_zona__in=zonas_coincidentes,
                id_tecnico__turno=turno_actual,
                activo=True
            )

            mejor_asignacion = None
            menor_carga = None

            for asignacion in asignaciones:
                tecnico = asignacion.id_tecnico

                pendientes = OrdenAtencion.objects.filter(
                    id_tecnico=tecnico,
                    estado="por_atender"
                ).exclude(id_orden=orden.id_orden).count()

                if menor_carga is None or pendientes < menor_carga:
                    menor_carga = pendientes
                    mejor_asignacion = asignacion

            if mejor_asignacion:
                tecnico_asignado = mejor_asignacion.id_tecnico
                zona_asignada = mejor_asignacion.id_zona
            elif zonas_coincidentes:
                zona_asignada = zonas_coincidentes[0]

        orden.id_tecnico = tecnico_asignado
        orden.id_zona = zona_asignada

    orden.save()

    return redirect("panel_control_interno")




